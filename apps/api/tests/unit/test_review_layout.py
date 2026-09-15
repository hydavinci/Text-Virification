import base64
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree

import fitz
import pytest
from docx import Document

from text_verification.application import original_preview


def test_short_table_fields_and_repeated_values_map_by_lines_and_neighboring_labels() -> None:
    from text_verification.application.review_layout import render_review_document

    body = "The agreement with Acme Company preserves the original long body paragraph."
    source = f"Title\n{body}\nAcme Company\nFee\n7380\nOther\nTotal\n7380\nEnd"
    rendered = f"Title\nAcme Company\nTotal\n7380\nEnd\nFee\n7380\nOther\n{body}"
    with fitz.open() as pdf:
        pdf.new_page().insert_text((30, 100), rendered)
        content = pdf.tobytes()

    layout = render_review_document(content, "pdf", source, source)

    mapped = {g.start: g for page in layout.pages for g in page.glyphs}
    assert set(mapped) == {i for i, char in enumerate(source) if char != "\n"}
    assert mapped[source.index("7380")].y > mapped[source.rindex("7380")].y
    assert mapped[source.rindex("Acme Company")].y < mapped[source.index("The agreement")].y
    assert layout.notice is None


def test_short_wrapped_field_matches_only_complete_rendered_lines() -> None:
    from text_verification.application.review_layout import render_review_document

    body = "The original body paragraph remains fully available."
    source = f"{body}\nTotal:"
    with fitz.open() as pdf:
        pdf.new_page().insert_text((30, 100), f"To\ntal:\n{body}")
        content = pdf.tobytes()

    layout = render_review_document(content, "pdf", source, source)

    assert {g.start for page in layout.pages for g in page.glyphs} == {
        i for i, char in enumerate(source) if char != "\n"
    }
    assert layout.notice is None


@pytest.mark.parametrize("marker_font", ["zapfdingbats", "helv"])
def test_short_field_can_omit_a_symbol_font_marker_but_not_regular_text(
    marker_font: str,
) -> None:
    from text_verification.application.review_layout import render_review_document

    body = "The original body paragraph remains fully available."
    source = f"{body}\nPro"
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((30, 100), "R", fontname=marker_font)
        page.insert_text((40, 100), "Pro")
        page.insert_text((30, 150), body)
        content = pdf.tobytes()

    layout = render_review_document(content, "pdf", source, source)

    marks = [g for page in layout.pages for g in page.glyphs if g.start >= source.index("Pro")]
    if marker_font == "zapfdingbats":
        assert {g.start for g in marks} == set(range(source.index("Pro"), len(source)))
        assert all(g.x >= 40 for g in marks)
        assert layout.notice is None
    else:
        assert not marks
        assert layout.notice


@pytest.mark.parametrize(
    ("source", "rendered", "unmapped"),
    [
        ("Fee\nCompletely unavailable source paragraph", "ServiceFee\nOther content", "Fee"),
        ("Heading\nHeading plus missing body", "Heading\nplus different body", "Heading"),
        ("Fee\nKnown original long body", "Fee\nFee\nKnown original long body", "Fee"),
        (
            "Fee\n7380\nCompletely unavailable middle text\nTotal\n7380",
            "Total\n7380\nUnrelated middle text\nFee\n7380", "7380",
        ),
    ],
)
def test_short_field_fallback_keeps_unresolved_ambiguity_explicit(
    source: str, rendered: str, unmapped: str,
) -> None:
    from text_verification.application.review_layout import render_review_document

    with fitz.open() as pdf:
        pdf.new_page().insert_text((30, 100), rendered)
        content = pdf.tobytes()
    layout = render_review_document(content, "pdf", source, source)
    start = source.index(unmapped)
    assert not any(
        g.start < start + len(unmapped) and g.end > start for p in layout.pages for g in p.glyphs
    )
    assert layout.notice


def test_overprinted_mixed_font_spaces_map_to_the_measured_anchor_gap() -> None:
    from text_verification.application.review_layout import _Glyph, _map_glyphs
    from text_verification.domain.review_layout import LayoutPage

    source = "A  5   B"
    # Office underlines can overprint spaces and split one visual line into
    # several raw text lines; the numeric run also has different font metrics.
    glyphs = [
        _Glyph("A", 0, 247.600006, 453.370789, 10.5, 15.203979, 30),
        _Glyph(" ", 0, 258.100006, 453.370789, 2.35199, 15.203979, 30),
        _Glyph(" ", 0, 260.451996, 453.370789, 2.35199, 15.203979, 30),
        _Glyph(" ", 0, 258.0, 453.370789, 2.35199, 15.203979, 30),
        _Glyph(" ", 0, 260.35199, 453.370789, 2.35199, 15.203979, 30),
        _Glyph("5", 0, 262.799988, 455.806793, 6.678009, 12.211487, 30),
        _Glyph(" ", 0, 269.446503, 455.806793, 3.328491, 12.211487, 30),
        _Glyph(" ", 0, 272.785492, 455.806793, 3.328491, 12.211487, 30),
        _Glyph(" ", 0, 276.124481, 455.806793, 3.328491, 12.211487, 30),
        _Glyph(" ", 0, 262.700012, 453.370789, 2.35199, 15.203979, 31),
        _Glyph(" ", 0, 276.990509, 453.370789, 2.35199, 15.203979, 32),
        _Glyph("B", 0, 279.450012, 453.370789, 10.5, 15.203979, 32),
    ]
    page = LayoutPage(width=600, height=800, image="data:image/svg+xml;base64,", text=source)

    assert _map_glyphs(source, glyphs, [page])

    for start, end, left, right in [
        (1, 3, 258.100006, 262.799988),
        (4, 7, 269.477997, 279.450012),
    ]:
        spaces = [glyph for glyph in page.glyphs if glyph.start < end and glyph.end > start]
        assert {offset for glyph in spaces for offset in range(glyph.start, glyph.end)} == (
            set(range(start, end))
        )
        assert all(left - 0.0001 <= glyph.x < glyph.x + glyph.width <= right + 0.0001
                   for glyph in spaces)
        assert all(453 <= glyph.y <= 456 for glyph in spaces)


@pytest.mark.parametrize(
    ("source", "rendered"),
    [
        ("Alpha  beta", "Alpha  beta"),
        ("  Alpha", "  Alpha"),
        ("Alpha  ", "Alpha  "),
        ("Alpha  beta", "Alpha beta"),
        ("Alpha  ", "Alpha "),
        ("Alpha  \n  Beta", "Alpha \n  Beta"),
        ("Alpha\u00a0\u00a0beta", "Alpha  beta"),
    ],
)
def test_horizontal_whitespace_retains_its_anchored_page_coordinates(
    source: str, rendered: str,
) -> None:
    from text_verification.application.review_layout import render_review_document

    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((30, 100), rendered)
        content = pdf.tobytes()
    layout = render_review_document(content, "pdf", source, source)
    glyphs = layout.pages[0].glyphs
    assert {index for glyph in glyphs for index in range(glyph.start, glyph.end)} == {
        index for index, character in enumerate(source) if character != "\n"
    }
    assert [glyph.start for glyph in glyphs] == sorted(glyph.start for glyph in glyphs)
    assert all(glyph.width > 0 and glyph.height > 0 for glyph in glyphs)
    assert layout.notice is None


@pytest.mark.parametrize(
    ("source", "rendered"),
    [
        ("Alpha  ", "Alpha extra"),
        ("  Alpha", "prefix Alpha"),
        ("StartingAlpha  EndingBeta", "StartingAlpha EXTRA EndingBeta"),
    ],
)
def test_whitespace_does_not_borrow_coordinates_from_unmatched_text(
    source: str, rendered: str,
) -> None:
    from text_verification.application.review_layout import render_review_document

    with fitz.open() as pdf:
        pdf.new_page().insert_text((30, 100), rendered)
        content = pdf.tobytes()
    layout = render_review_document(content, "pdf", source, source)
    assert any(page.glyphs for page in layout.pages)
    assert all(
        not source[glyph.start:glyph.end].isspace()
        for page in layout.pages for glyph in page.glyphs
    )


def test_unique_paragraphs_map_when_table_extraction_changes_reading_order() -> None:
    from text_verification.application.review_layout import render_review_document

    source = "Original content\nPlease  confirm payment.\nTable preserved"
    rendered = "Original content\nTable preserved\nPlease  confirm payment."
    with fitz.open() as pdf:
        pdf.new_page().insert_text((30, 100), rendered)
        content = pdf.tobytes()
    layout = render_review_document(content, "pdf", source, source)
    glyphs = layout.pages[0].glyphs
    assert {glyph.start for glyph in glyphs} == {
        index for index, char in enumerate(source) if char != "\n"
    }
    table = next(glyph for glyph in glyphs if glyph.start == source.index("Table"))
    body = next(glyph for glyph in glyphs if glyph.start == source.index("Please"))
    assert table.y < body.y
    assert layout.notice is None


def test_unique_paragraph_stays_locatable_without_guessing_unmatched_paragraphs() -> None:
    from text_verification.application.review_layout import render_review_document

    source = "Known exact paragraph\nCompletely unavailable text that must not be guessed"
    with fitz.open() as pdf:
        pdf.new_page().insert_text((30, 100), "Known exact paragraph\nOther content")
        content = pdf.tobytes()
    layout = render_review_document(content, "pdf", source, source)
    assert {glyph.start for glyph in layout.pages[0].glyphs} == set(range(21))
    assert layout.notice


@pytest.mark.parametrize("body_changed", [False, True])
def test_paragraph_fallback_does_not_reuse_body_coordinates_for_a_missing_heading(
    body_changed: bool,
) -> None:
    from text_verification.application.review_layout import render_review_document

    heading = "Preserved heading"
    body = f"{heading} plus unique body"
    source = f"{heading}\n{body}\nCompletely unavailable paragraph"
    rendered_body = body.replace("unique", "revised") if body_changed else body
    with fitz.open() as pdf:
        pdf.new_page().insert_text((30, 100), f"{rendered_body}\nOther content")
        content = pdf.tobytes()
    layout = render_review_document(content, "pdf", source, source)
    body_start = len(heading) + 1
    expected = set() if body_changed else set(range(body_start, body_start + len(body)))
    assert {glyph.start for glyph in layout.pages[0].glyphs} == expected
    assert layout.notice


def test_raw_order_spaces_on_another_line_are_not_used_as_coordinates() -> None:
    from text_verification.application.review_layout import render_review_document

    source = "Alpha  beta"
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((30, 100), "Alpha")
        page.insert_text((58.13, 200), "  ")
        page.insert_text((64.25, 100), "beta")
        content = pdf.tobytes()
    layout = render_review_document(content, "pdf", source, source)
    assert not any(glyph.start in {5, 6} for glyph in layout.pages[0].glyphs)


@pytest.mark.parametrize("rotation", [90, 180, 270])
def test_rotated_whitespace_retains_its_actual_coordinates(rotation: int) -> None:
    from text_verification.application.review_layout import render_review_document

    source = "Alpha  beta"
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((30, 100), source)
        page.set_rotation(rotation)
        content = pdf.tobytes()
    layout = render_review_document(content, "pdf", source, source)
    assert {glyph.start for glyph in layout.pages[0].glyphs} == set(range(len(source)))


def sample_pdf(*, rotated: bool = False) -> bytes:
    with fitz.open() as pdf:
        for word in ("First account", "Second account"):
            page = pdf.new_page(width=300, height=400)
            page.insert_text((20, 25), "Repeated header")
            page.insert_text((30, 100), word)
            if rotated:
                page.set_rotation(90)
        return pdf.tobytes()


def test_layout_maps_repeated_words_to_their_actual_pages() -> None:
    from text_verification.application.review_layout import render_review_document

    text = "First account\nSecond account"
    layout = render_review_document(sample_pdf(), "pdf", text, text)
    assert len(layout.pages) == 2
    assert layout.revision_applied
    assert layout.pages[0].image.startswith("data:image/svg+xml;base64,")
    for page, word in zip(layout.pages, ("First account", "Second account"), strict=True):
        start = text.index(word)
        assert {g.start for g in page.glyphs} == set(range(start, start + len(word)))
        assert all(0 <= g.x < g.x + g.width <= page.width for g in page.glyphs)
        assert all(0 <= g.y < g.y + g.height <= page.height for g in page.glyphs)


def test_page_text_stays_vector_instead_of_a_low_resolution_page_bitmap() -> None:
    from text_verification.application.review_layout import render_review_document

    text = "First account\nSecond account"
    layout = render_review_document(sample_pdf(), "pdf", text, text)
    for page in layout.pages:
        assert page.image.startswith("data:image/svg+xml;base64,")
        svg = ElementTree.fromstring(base64.b64decode(page.image.split(",", 1)[1]))
        assert svg.tag == "{http://www.w3.org/2000/svg}svg"
        assert len(svg.findall(".//{http://www.w3.org/2000/svg}path")) >= 10
        assert svg.findall(".//{http://www.w3.org/2000/svg}use")
        assert not svg.findall(".//{http://www.w3.org/2000/svg}image")
        assert not svg.findall(".//{http://www.w3.org/2000/svg}text")


def test_rotated_page_coordinates_follow_the_rendered_page() -> None:
    from text_verification.application.review_layout import render_review_document

    layout = render_review_document(
        sample_pdf(rotated=True), "pdf", "First account", "First account",
    )
    page = layout.pages[0]
    assert page.width == 400
    assert page.height == 300
    assert page.glyphs[0].x > 250
    svg = ElementTree.fromstring(base64.b64decode(page.image.split(",", 1)[1]))
    assert list(map(float, svg.attrib["viewBox"].split())) == [0, 0, page.width, page.height]


def test_word_preview_applies_revisions_without_touching_the_source(tmp_path: Path) -> None:
    from text_verification.application.review_layout import render_review_document

    document = Document()
    document.sections[0].header.paragraphs[0].text = "Header preserved"
    document.add_paragraph("First account")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Second account"
    image = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
    image.clear_with(125)
    document.add_picture(BytesIO(image.tobytes("png")))
    path = tmp_path / "original.docx"
    document.save(path)
    source = path.read_bytes()
    text = "First account\nSecond account\nHeader preserved"

    def convert(content: bytes, file_type: str, target_type: str) -> bytes:
        assert (file_type, target_type) == ("docx", "pdf")
        revised = Document(BytesIO(content))
        assert revised.paragraphs[0].text == "First corrected"
        assert len(revised.inline_shapes) == 1
        assert revised.tables[0].cell(0, 0).text == "Second account"
        assert revised.sections[0].header.paragraphs[0].text == "Header preserved"
        with fitz.open() as pdf:
            page = pdf.new_page()
            page.insert_text((30, 60), "First corrected\nSecond account")
            return pdf.tobytes()

    layout = render_review_document(
        source, "docx", text, text.replace("First account", "First corrected"), convert=convert,
    )
    assert layout.revision_applied
    assert "corrected" in layout.pages[0].text
    assert path.read_bytes() == source


def test_image_revision_is_explicitly_an_annotation_not_replaced_pixels() -> None:
    from text_verification.application.review_layout import render_review_document

    image = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 100, 100), False)
    image.clear_with(255)
    layout = render_review_document(image.tobytes("png"), "png", "wrong", "right")
    assert not layout.revision_applied
    assert layout.notice
    assert len(layout.pages) == 1


def test_mapping_does_not_guess_coordinates_for_unrelated_text() -> None:
    from text_verification.application.review_layout import render_review_document

    layout = render_review_document(sample_pdf(), "pdf", "unrelated", "unrelated")
    assert all(not page.glyphs for page in layout.pages)
    assert layout.notice


def test_mapping_does_not_choose_between_identical_header_and_body() -> None:
    from text_verification.application.review_layout import render_review_document

    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((20, 25), "Account")
        page.insert_text((20, 100), "Account")
        content = pdf.tobytes()
    layout = render_review_document(content, "pdf", "Account", "Account")
    assert not layout.pages[0].glyphs
    assert layout.notice


def test_heading_repeated_inside_body_is_not_mistaken_for_an_extra_header() -> None:
    from text_verification.application.review_layout import render_review_document

    heading = "Project delivery report"
    text = f"{heading}\nThe {heading} is ready for approval."
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((30, 100), text)
        content = pdf.tobytes()

    layout = render_review_document(content, "pdf", text, text)
    assert {glyph.start for page in layout.pages for glyph in page.glyphs} == {
        index for index, character in enumerate(text) if character != "\n"
    }
    assert layout.notice is None


def test_extra_header_remains_ambiguous_when_heading_also_occurs_inside_body() -> None:
    from text_verification.application.review_layout import render_review_document

    heading = "Project delivery report"
    text = f"{heading}\nThe {heading} is ready for approval."
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((30, 25), heading)
        page.insert_text((30, 100), text)
        content = pdf.tobytes()

    layout = render_review_document(content, "pdf", text, text)
    assert not any(glyph.start < len(heading) for page in layout.pages for glyph in page.glyphs)
    assert any(glyph.start >= len(heading) for page in layout.pages for glyph in page.glyphs)
    assert layout.notice


def test_repeated_header_cannot_complete_a_false_body_match_across_pages() -> None:
    from text_verification.application.review_layout import render_review_document

    with fitz.open() as pdf:
        for body in ("Account summary", "Second account"):
            page = pdf.new_page()
            page.insert_text((20, 25), "Account summary")
            page.insert_text((20, 100), body)
        content = pdf.tobytes()
    text = "Account summary\nSecond account"
    layout = render_review_document(content, "pdf", text, text)
    assert not any(g.start < text.index("Second") for g in layout.pages[1].glyphs)
    assert layout.notice


def test_mapping_work_is_bounded(monkeypatch: pytest.MonkeyPatch) -> None:
    from text_verification.application import review_layout

    monkeypatch.setattr(review_layout, "MAX_MAPPING_WORK", 0)
    text = "First account\nSecond account"
    layout = review_layout.render_review_document(sample_pdf(), "pdf", text, text)
    assert all(not page.glyphs for page in layout.pages)
    assert layout.notice


@pytest.mark.parametrize("modified", [False, True])
def test_invalid_preview_is_reported_without_a_text_fallback(modified: bool) -> None:
    from text_verification.application.review_layout import render_review_document

    with pytest.raises(original_preview.OriginalPreviewError):
        render_review_document(b"broken", "pdf", "", "revision" if modified else "")
