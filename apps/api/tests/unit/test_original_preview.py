from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import fitz
import pytest
from docx import Document

from text_verification.domain.documents import DocumentModel, FileType
from text_verification.infrastructure.storage import InvalidUpload, JobStorage


def source_document(file_type: FileType, data: bytes) -> DocumentModel:
    return DocumentModel(
        document_id=uuid4(), source_version=f"sha256:{hashlib.sha256(data).hexdigest()}",
        file_type=file_type, source_name=f"source.{file_type.value}", text="", blocks=[],
        parser_name="test", parser_version="1",
    )


def test_pdf_preview_preserves_original_bytes(tmp_path: Path) -> None:
    from text_verification.application.original_preview import build_original_preview

    pdf = fitz.open()
    pdf.new_page().insert_text((30, 30), "Original page")
    data = pdf.tobytes()
    pdf.close()
    document = source_document(FileType.PDF, data)
    storage = JobStorage(tmp_path, 25 * 1024 * 1024)
    storage.save_bytes(document.document_id, document.source_name, data)

    preview = build_original_preview(document, storage, "http://unused")
    assert preview.media_type == "application/pdf"
    assert preview.content == data
    assert storage.source_path(document.document_id, FileType.PDF).read_bytes() == data


def test_word_preview_uses_verified_original_not_extracted_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    from text_verification.application import original_preview

    word = Document()
    word.add_paragraph("Keep original content")
    word.add_table(rows=2, cols=2).cell(0, 0).text = "Table cell"
    stream = BytesIO()
    word.save(stream)
    data = stream.getvalue()
    document = source_document(FileType.DOCX, data)
    storage = JobStorage(tmp_path, 25 * 1024 * 1024)
    storage.save_bytes(document.document_id, document.source_name, data)

    def render(content: bytes, file_type: str, renderer_url: str) -> bytes:
        assert content == data
        assert file_type == "docx"
        assert renderer_url == "http://renderer:8000"
        return b"%PDF-1.7\nrendered"

    monkeypatch.setattr(original_preview, "render_office", render)
    preview = original_preview.build_original_preview(document, storage, "http://renderer:8000")
    assert preview.content == b"%PDF-1.7\nrendered"
    assert storage.source_path(document.document_id, FileType.DOCX).read_bytes() == data


def test_changed_source_cannot_be_previewed(tmp_path: Path) -> None:
    from text_verification.application.original_preview import build_original_preview

    data = b"%PDF-1.7\noriginal"
    document = source_document(FileType.PDF, data)
    storage = JobStorage(tmp_path, 1024)
    stored = storage.save_bytes(document.document_id, document.source_name, data)
    stored.path.write_bytes(b"%PDF-1.7\nchanged")
    with pytest.raises(InvalidUpload):
        build_original_preview(document, storage, "http://unused")


def test_review_cache_reuses_revisions_but_always_verifies_source(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    from text_verification.application import original_preview

    data = b"%PDF-1.7\noriginal"
    document = source_document(FileType.PDF, data)
    storage = JobStorage(tmp_path, 1024)
    stored = storage.save_bytes(document.document_id, document.source_name, data)
    calls: list[bytes] = []

    def render(content: bytes, path: str, renderer_url: str, media_type: str) -> bytes:
        calls.append(content)
        return b'{"pages":[],"revision_applied":true,"notice":null}'

    monkeypatch.setattr(original_preview, "_request_renderer", render)
    first = original_preview.build_review_preview(document, storage, "http://renderer", "")
    again = original_preview.build_review_preview(document, storage, "http://renderer", "")
    assert first == again
    assert first is not again
    assert len(calls) == 1
    original_preview.build_review_preview(document, storage, "http://renderer", "revision")
    assert len(calls) == 2
    original_preview.build_review_preview(document, storage, "http://renderer", "")
    assert len(calls) == 2
    other = source_document(FileType.PDF, data)
    storage.save_bytes(other.document_id, other.source_name, data)
    original_preview.build_review_preview(other, storage, "http://renderer", "")
    assert len(calls) == 3
    original_preview.build_review_preview(document, storage, "http://different-renderer", "")
    assert len(calls) == 4
    stored.path.write_bytes(b"%PDF-1.7\nchanged")
    with pytest.raises(InvalidUpload):
        original_preview.build_review_preview(document, storage, "http://renderer", "")
    assert len(calls) == 4


def test_review_cache_does_not_retain_invalid_renderer_response(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    from text_verification.application import original_preview

    data = b"%PDF-1.7\noriginal"
    document = source_document(FileType.PDF, data)
    storage = JobStorage(tmp_path, 1024)
    storage.save_bytes(document.document_id, document.source_name, data)
    responses = iter([b"{}", b'{"pages":[],"revision_applied":true}'])

    def render(content: bytes, path: str, renderer_url: str, media_type: str) -> bytes:
        return next(responses)

    monkeypatch.setattr(original_preview, "_request_renderer", render)
    with pytest.raises(original_preview.OriginalPreviewError, match="无效"):
        original_preview.build_review_preview(document, storage, "http://renderer", "")
    assert original_preview.build_review_preview(
        document, storage, "http://renderer", "",
    ).revision_applied
    assert original_preview.build_review_preview(
        document, storage, "http://renderer", "",
    ).revision_applied
