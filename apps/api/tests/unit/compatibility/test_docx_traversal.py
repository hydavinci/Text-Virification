from copy import deepcopy
from io import BytesIO
from pathlib import Path

import fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches

from text_verification.application.original_preview import OriginalPreviewError
from text_verification.application.review_layout import render_review_document
from text_verification.compatibility import parser
from text_verification.compatibility.exporters import ExportError, export_original


@pytest.fixture
def structured_docx(tmp_path: Path) -> Path:
    document = Document()
    document.add_paragraph("Body start")
    cell = document.add_table(rows=1, cols=2).cell(0, 0)
    cell.text = "Cell before"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested"
    cell.add_paragraph("Cell after")
    document.tables[0].cell(0, 1).text = "Other cell"
    document.add_paragraph("Body middle")
    merged = document.add_table(rows=2, cols=2)
    merged.cell(0, 0).merge(merged.cell(1, 1)).text = "Merged"
    document.add_paragraph("Body end")
    section = document.sections[0]
    section.header.paragraphs[0].text = "Header"
    section.header.add_table(rows=1, cols=1, width=Inches(1)).cell(0, 0).text = "Header cell"
    section.footer.paragraphs[0].text = "Footer"
    section.first_page_header.paragraphs[0].text = "First header"
    section.even_page_footer.paragraphs[0].text = "Even footer"
    document.add_section()  # The same header/footer parts must not be extracted again.
    path = tmp_path / "structured.docx"
    document.save(path)
    return path


EXPECTED = (
    "Body start\nCell before\nNested\nCell after\nOther cell\nBody middle\nMerged\n"
    "Body end\nHeader\nHeader cell\nFirst header\nFooter\nEven footer"
)


def test_extracts_interleaved_nested_and_shared_stories_once(structured_docx: Path) -> None:
    text, page_map = parser._parse_docx(str(structured_docx))

    assert text == EXPECTED
    assert [text[start:end] for start, end, _ in page_map] == EXPECTED.splitlines()
    assert [label for _, _, label in page_map] == [
        f"第{index}段" for index in range(1, 14)
    ]
    assert all(
        next_start == end + 1
        for (_, end, _), (next_start, _, _) in zip(page_map, page_map[1:], strict=False)
    )


@pytest.mark.parametrize("target", ["Nested", "Merged", "Body end", "Header cell", "Even footer"])
@pytest.mark.parametrize("tracked", [False, True])
def test_export_changes_only_the_positioned_source_paragraph(
    structured_docx: Path, tmp_path: Path, target: str, tracked: bool,
) -> None:
    before = structured_docx.read_bytes()
    start = EXPECTED.index(target)
    result = export_original(
        structured_docx, "docx", [(target, "Revised", start, start + len(target))],
        tracked, original_text=EXPECTED,
    )
    restored = Document(BytesIO(result.content))
    if tracked:
        # Accept revisions locally to exercise the same parser on the resulting package.
        for part in restored.part.package.parts:
            if not hasattr(part, "element"):
                continue
            for deletion in part.element.xpath(".//w:del"):
                deletion.getparent().remove(deletion)
            for insertion in part.element.xpath(".//w:ins"):
                parent = insertion.getparent()
                index = parent.index(insertion)
                for child in list(insertion):
                    parent.insert(index, child)
                    index += 1
                parent.remove(insertion)
    output = tmp_path / "revised.docx"
    restored.save(output)
    assert parser._parse_docx(str(output))[0] == EXPECTED.replace(target, "Revised")
    assert structured_docx.read_bytes() == before


def test_repeated_text_is_edited_by_position_not_first_match(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Repeated")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Repeated"
    document.add_paragraph("End")
    path = tmp_path / "repeated.docx"
    document.save(path)
    text = "Repeated\nRepeated\nEnd"
    result = export_original(
        path, "docx", [("Repeated", "Changed", 9, 17)], False, original_text=text,
    )
    restored = Document(BytesIO(result.content))
    assert restored.paragraphs[0].text == "Repeated"
    assert restored.tables[0].cell(0, 0).text == "Changed"
    assert restored.paragraphs[1].text == "End"


@pytest.mark.parametrize("same_text", [False, True])
def test_legacy_source_order_never_silently_retargets_an_edit(
    tmp_path: Path, same_text: bool,
) -> None:
    document = Document()
    document.add_paragraph("Start")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Same" if same_text else "Cell"
    document.add_paragraph("Same" if same_text else "Tail")
    path = tmp_path / "legacy.docx"
    document.save(path)
    old_text = "Start\nSame\nSame" if same_text else "Start\nTail\nCell"
    with pytest.raises(ExportError, match="source|order|reanaly"):
        export_original(
            path, "docx", [(old_text[6:10], "Edit", 6, 10)],
            False, original_text=old_text,
        )


@pytest.mark.parametrize("kind", ["markup", "hyperlink"])
def test_non_linear_run_mapping_is_rejected_without_shifting_later_paragraphs(
    tmp_path: Path, kind: str,
) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    if kind == "markup":
        paragraph.add_run("<b>Linked</b>")
    else:
        hyperlink = OxmlElement("w:hyperlink")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Linked"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
    document.add_paragraph("Later")
    path = tmp_path / "inline.docx"
    document.save(path)
    text = "Linked\nLater"
    result = export_original(
        path, "docx", [("Later", "After", 7, 12)], False, original_text=text,
    )
    assert Document(BytesIO(result.content)).paragraphs[1].text == "After"
    with pytest.raises(ExportError, match="safely|mapping"):
        export_original(
            path, "docx", [("Linked", "Edited", 0, 6)], False, original_text=text,
        )


@pytest.mark.parametrize("limit", ["MAX_DOC_PARSED_ELEMENTS", "MAX_DOC_PARSED_TEXT_CHARS"])
@pytest.mark.parametrize("story", ["nested", "header"])
def test_newly_covered_text_obeys_existing_resource_limits(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, limit: str, story: str,
) -> None:
    document = Document()
    if story == "nested":
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        container = cell.add_table(rows=1, cols=1).cell(0, 0)
    else:
        container = document.sections[0].header
    container.paragraphs[0].text = "123456"
    container.add_paragraph("789")
    path = tmp_path / "limited.docx"
    document.save(path)
    monkeypatch.setattr(parser, limit, 1)
    with pytest.raises(ValueError, match="limit"):
        parser._parse_docx(str(path))


@pytest.mark.parametrize("modified", [False, True])
def test_preview_rejects_legacy_source_order_before_rendering(
    tmp_path: Path, modified: bool,
) -> None:
    document = Document()
    document.add_paragraph("Start")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Cell"
    document.add_paragraph("Tail")
    path = tmp_path / "legacy.docx"
    document.save(path)

    def convert(*args: object) -> bytes:
        pytest.fail("unsafe source mapping must be rejected before rendering")

    with pytest.raises(OriginalPreviewError):
        render_review_document(
            path.read_bytes(), "docx", "Start\nTail\nCell",
            "Start\nEdit\nCell" if modified else "Start\nTail\nCell", convert=convert,
        )


def test_preview_applies_nested_revision_using_the_canonical_order(
    structured_docx: Path, tmp_path: Path,
) -> None:
    revised = EXPECTED.replace("Nested", "Edited")

    def convert(content: bytes, source_type: str, target_type: str) -> bytes:
        assert (source_type, target_type) == ("docx", "pdf")
        path = tmp_path / "preview.docx"
        path.write_bytes(content)
        assert parser._parse_docx(str(path))[0] == revised
        with fitz.open() as pdf:
            pdf.new_page().insert_text((30, 50), revised)
            return pdf.tobytes()

    layout = render_review_document(
        structured_docx.read_bytes(), "docx", EXPECTED, revised, convert=convert,
    )
    assert layout.revision_applied
    assert layout.notice is None
    assert any(g.start == revised.index("Edited") for p in layout.pages for g in p.glyphs)


@pytest.mark.parametrize("render_body", [False, True])
def test_preview_does_not_swap_identical_body_and_header_locations(
    tmp_path: Path, render_body: bool,
) -> None:
    document = Document()
    document.add_paragraph("Account")
    document.sections[0].header.paragraphs[0].text = "Account"
    path = tmp_path / "ambiguous.docx"
    document.save(path)

    def convert(content: bytes, source_type: str, target_type: str) -> bytes:
        with fitz.open() as pdf:
            page = pdf.new_page()
            page.insert_text((30, 25), "Account")  # Header comes first in the PDF.
            if render_body:
                page.insert_text((30, 100), "Account")
            return pdf.tobytes()

    layout = render_review_document(
        path.read_bytes(), "docx", "Account\nAccount", "Account\nAccount", convert=convert,
    )
    assert not any(page.glyphs for page in layout.pages)
    assert layout.notice


@pytest.mark.parametrize(
    ("body_paragraphs", "header_paragraphs"),
    [
        (["Account\nNumber"], ["Account", "Number"]),
        (["Account", "Number"], ["Account\nNumber"]),
        (["Acc", "ountNumber"], ["AccountNum", "ber"]),
    ],
)
@pytest.mark.parametrize("unique_text", [False, True])
def test_preview_requires_story_evidence_independent_of_paragraph_boundaries(
    tmp_path: Path,
    body_paragraphs: list[str],
    header_paragraphs: list[str],
    unique_text: bool,
) -> None:
    document = Document()
    for text in body_paragraphs:
        document.add_paragraph(text)
    header = document.sections[0].header
    header.paragraphs[0].text = header_paragraphs[0]
    for text in header_paragraphs[1:]:
        header.add_paragraph(text)
    if unique_text:
        document.add_paragraph("Unique body evidence")
        header.add_paragraph("Unique header evidence")
    path = tmp_path / "segmentation.docx"
    document.save(path)
    body = "\n".join(body_paragraphs)
    heading = "\n".join(header_paragraphs)
    if unique_text:
        body += "\nUnique body evidence"
        heading += "\nUnique header evidence"
    source = f"{body}\n{heading}"
    assert parser._parse_docx(str(path))[0] == source

    def convert(content: bytes, source_type: str, target_type: str) -> bytes:
        with fitz.open() as pdf:
            page = pdf.new_page()
            page.insert_text((30, 25), heading)
            page.insert_text((30, 150), body)
            return pdf.tobytes()

    layout = render_review_document(
        path.read_bytes(), "docx", source, source, convert=convert,
    )
    mapped = {g.start: g for page in layout.pages for g in page.glyphs}
    expected = set()
    if unique_text:
        for text in ("Unique body evidence", "Unique header evidence"):
            expected.update(range(source.index(text), source.index(text) + len(text)))
    assert set(mapped) == expected
    assert layout.notice
    if unique_text:
        assert mapped[source.index("Unique body evidence")].y > 150
        assert mapped[source.index("Unique header evidence")].y < 100


def test_distinct_header_parts_are_not_deduplicated_by_their_text(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Body")
    document.sections[0].header.paragraphs[0].text = "Repeated header"
    section = document.add_section()
    section.header.is_linked_to_previous = False
    section.header.paragraphs[0].text = "Repeated header"
    # Explicitly reference the first footer part from both sections.
    document.sections[0].footer.paragraphs[0].text = "Shared footer"
    section._sectPr.append(deepcopy(document.sections[0]._sectPr.footerReference_lst[0]))
    path = tmp_path / "parts.docx"
    document.save(path)
    before = path.read_bytes()

    assert parser._parse_docx(str(path))[0] == (
        "Body\nRepeated header\nShared footer\nRepeated header"
    )
    assert path.read_bytes() == before


def test_plain_docx_traversal_does_not_create_header_footer_parts(tmp_path: Path) -> None:
    from text_verification.compatibility.docx_traversal import iter_docx_paragraphs

    document = Document()
    document.add_paragraph("Body")
    before = len(document.part.package.parts)
    assert [paragraph.text for paragraph in iter_docx_paragraphs(document)] == ["Body"]
    assert len(document.part.package.parts) == before


@pytest.mark.parametrize("replacement", ["", " Padded "])
@pytest.mark.parametrize("tracked", [False, True])
def test_export_rejects_revisions_that_cannot_roundtrip_extracted_offsets(
    tmp_path: Path, replacement: str, tracked: bool,
) -> None:
    document = Document()
    for text in ("Start", "Target", "End"):
        document.add_paragraph(text)
    path = tmp_path / "normalization.docx"
    document.save(path)
    with pytest.raises(ExportError, match="mapping|extraction|paragraph"):
        export_original(
            path, "docx", [("Target", replacement, 6, 12)], tracked,
            original_text="Start\nTarget\nEnd",
        )
