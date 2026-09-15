from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from text_verification.compatibility import exporters
from text_verification.compatibility.exporters import export_original


def test_libreoffice_core_metadata_is_preserved_without_duplicate_archive_parts(
    tmp_path: Path,
) -> None:
    document = Document()
    document.add_paragraph("cat test@example.com")
    document.core_properties.author = "Example author"
    document.core_properties.title = "Example title"
    stream = BytesIO()
    document.save(stream)
    source = tmp_path / "libreoffice.docx"
    alternate_type = (
        "http://schemas.openxmlformats.org/officedocument/2006/"
        "relationships/metadata/core-properties"
    )
    with ZipFile(BytesIO(stream.getvalue())) as original, ZipFile(source, "w") as output:
        for name in original.namelist():
            content = original.read(name)
            if name == "_rels/.rels":
                content = content.replace(RT.CORE_PROPERTIES.encode(), alternate_type.encode())
            output.writestr(name, content)

    exported = export_original(
        source, "docx", [], False,
        original_text="cat test@example.com", modified_text="dog test@example.com",
    )

    with ZipFile(BytesIO(exported.content)) as archive:
        assert len(archive.namelist()) == len(set(archive.namelist()))
    restored = Document(BytesIO(exported.content))
    assert restored.paragraphs[0].text == "dog test@example.com"
    assert restored.core_properties.author == "Example author"
    assert restored.core_properties.title == "Example title"


def test_tracked_docx_export_is_byte_deterministic_across_wall_clock_changes(
    tmp_path: Path,
    monkeypatch,
) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("帐号测试")
    document.save(source)
    times = iter(
        [
            datetime(2026, 9, 4, 1, 0, tzinfo=UTC),
            datetime(2026, 9, 4, 2, 0, tzinfo=UTC),
        ]
    )

    class ChangingDateTime:
        @classmethod
        def now(cls, tz=None):
            del cls, tz
            return next(times)

    monkeypatch.setattr(
        exporters,
        "datetime",
        ChangingDateTime,
        raising=False,
    )

    first = export_original(
        source,
        "docx",
        [],
        True,
        original_text="帐号测试",
        modified_text="账号测试",
    )
    second = export_original(
        source,
        "docx",
        [],
        True,
        original_text="帐号测试",
        modified_text="账号测试",
    )

    assert first.content == second.content
