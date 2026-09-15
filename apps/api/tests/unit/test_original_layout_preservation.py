from __future__ import annotations

import base64
import hashlib
import os
import re
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

import fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from text_verification.application.original_preview import build_review_preview, render_office
from text_verification.compatibility.exporters import export_original
from text_verification.compatibility.parser import _parse_docx
from text_verification.domain.documents import DocumentModel
from text_verification.infrastructure.storage import JobStorage


@pytest.fixture
def illustrated_word(tmp_path: Path) -> Path:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Header preserved"
    document.sections[0].footer.paragraphs[0].text = "Footer preserved"
    document.add_paragraph("Original content")
    document.add_table(rows=2, cols=2).cell(0, 0).text = "Table preserved"
    image = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 32, 32), False)
    image.clear_with(120)
    document.add_picture(BytesIO(image.tobytes("png")), width=Inches(1))
    path = tmp_path / "illustrated.docx"
    document.save(path)
    return path


@pytest.mark.parametrize("track_changes", [False, True])
def test_word_correction_keeps_images_tables_headers_and_footers(
    illustrated_word: Path, track_changes: bool,
) -> None:
    original_bytes = illustrated_word.read_bytes()
    text, _ = _parse_docx(str(illustrated_word))
    start = text.index("Original")
    exported = export_original(
        illustrated_word, "docx", [("Original", "Revised", start, start + 8)],
        track_changes, original_text=text,
    )
    restored = Document(BytesIO(exported.content))
    assert len(restored.inline_shapes) == 1
    assert restored.tables[0].cell(0, 0).text == "Table preserved"
    assert restored.sections[0].header.paragraphs[0].text == "Header preserved"
    assert restored.sections[0].footer.paragraphs[0].text == "Footer preserved"
    with ZipFile(BytesIO(original_bytes)) as original, ZipFile(BytesIO(exported.content)) as edited:
        media = [name for name in original.namelist() if name.startswith("word/media/")]
        assert media
        assert all(original.read(name) == edited.read(name) for name in media)
        assert b"Revised" in edited.read("word/document.xml")
    assert illustrated_word.read_bytes() == original_bytes


def test_docker_renderer_preserves_illustrated_word_pages(illustrated_word: Path) -> None:
    url = os.environ.get("PREVIEW_RENDERER_TEST_URL")
    if not url:
        pytest.skip("Set PREVIEW_RENDERER_TEST_URL to exercise the Docker renderer.")
    rendered = render_office(illustrated_word.read_bytes(), "docx", url)
    with fitz.open(stream=rendered, filetype="pdf") as pdf:
        assert len(pdf) >= 1
        text = "\n".join(page.get_text() for page in pdf)
        for expected in (
            "Header preserved", "Footer preserved", "Original content", "Table preserved",
        ):
            assert expected in text
        assert any(page.get_images() for page in pdf)


def test_docker_renderer_survives_repeated_office_conversions(illustrated_word: Path) -> None:
    url = os.environ.get("PREVIEW_RENDERER_TEST_URL")
    if not url:
        pytest.skip("Set PREVIEW_RENDERER_TEST_URL to exercise the Docker renderer.")
    content = illustrated_word.read_bytes()
    # Office spawns several helpers per conversion; unreaped children exhaust the 128-PID limit.
    for _ in range(30):
        assert render_office(content, "docx", url).startswith(b"%PDF-")


def test_docker_unified_review_renders_revisions_and_preserves_the_picture(
    illustrated_word: Path, tmp_path: Path,
) -> None:
    url = os.environ.get("PREVIEW_RENDERER_TEST_URL")
    if not url:
        pytest.skip("Set PREVIEW_RENDERER_TEST_URL to exercise the Docker renderer.")
    content = illustrated_word.read_bytes()
    text, _ = _parse_docx(str(illustrated_word))
    job_id = uuid4()
    document = DocumentModel(
        document_id=job_id, source_name="illustrated.docx", file_type="docx",
        source_version=f"sha256:{hashlib.sha256(content).hexdigest()}",
        text=text, blocks=[], parser_name="test", parser_version="1",
    )
    storage = JobStorage(tmp_path / "jobs", 25 * 1024 * 1024)
    source = storage.save_bytes(job_id, "illustrated.docx", content)
    revised = text.replace("Original", "Revised")
    layout = build_review_preview(document, storage, url, revised)
    assert layout.revision_applied
    preview_text = "\n".join(page.text for page in layout.pages)
    assert "Revised content" in preview_text
    assert "Original content" not in preview_text
    for expected in ("Table preserved", "Header preserved", "Footer preserved"):
        assert expected in preview_text
    assert any(g.start == revised.index("Revised") for page in layout.pages for g in page.glyphs)
    gray_pixels = 0
    for page in layout.pages:
        assert page.image.startswith("data:image/svg+xml;base64,")
        with fitz.open(stream=base64.b64decode(page.image.split(",", 1)[1]), filetype="svg") as svg:
            gray_pixels += svg[0].get_pixmap(alpha=False).samples.count(bytes([120] * 3))
    assert gray_pixels > 1000
    assert source.path.read_bytes() == content


def test_docker_unified_review_maps_spaces_in_word_paragraphs(
    illustrated_word: Path, tmp_path: Path,
) -> None:
    url = os.environ.get("PREVIEW_RENDERER_TEST_URL")
    if not url:
        pytest.skip("Set PREVIEW_RENDERER_TEST_URL to exercise the Docker renderer.")
    word = Document(illustrated_word)
    word.add_paragraph("Please  confirm the payment.")
    word.add_paragraph("乙方应在  5 个工作日内提出。")
    word.save(illustrated_word)
    content = illustrated_word.read_bytes()
    text, _ = _parse_docx(str(illustrated_word))
    job_id = uuid4()
    document = DocumentModel(
        document_id=job_id, source_name="spaces.docx", file_type="docx",
        source_version=f"sha256:{hashlib.sha256(content).hexdigest()}",
        text=text, blocks=[], parser_name="test", parser_version="1",
    )
    storage = JobStorage(tmp_path / "jobs", 25 * 1024 * 1024)
    source = storage.save_bytes(job_id, "spaces.docx", content)
    for revision in (text, text.replace("Please  confirm", "Please confirm")):
        layout = build_review_preview(document, storage, url, revision)
        runs = list(re.finditer(" {2,}", revision))
        assert runs
        for run in runs:
            mapped = {
                index for page in layout.pages for glyph in page.glyphs
                for index in range(max(run.start(), glyph.start), min(run.end(), glyph.end))
            }
            assert mapped == set(range(run.start(), run.end()))
        assert source.path.read_bytes() == content


def test_docker_unified_review_maps_short_table_fields_and_a_checkbox(tmp_path: Path) -> None:
    url = os.environ.get("PREVIEW_RENDERER_TEST_URL")
    if not url:
        pytest.skip("Set PREVIEW_RENDERER_TEST_URL to exercise the Docker renderer.")
    word = Document()
    word.add_paragraph("Title")
    word.add_paragraph("The agreement with Acme Company preserves the original body paragraph.")
    table = word.add_table(rows=4, cols=2)
    for row, values in zip(table.rows, [
        ("Vendor", "Acme Company"), ("Fee", "7380"), ("Total", "7380"), ("Level", ""),
    ], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    paragraph = table.cell(3, 1).paragraphs[0]
    symbol = OxmlElement("w:sym")
    symbol.set(qn("w:font"), "Wingdings 2")
    symbol.set(qn("w:char"), "F052")
    paragraph.add_run()._r.append(symbol)
    paragraph.add_run("Pro")
    path = tmp_path / "short-fields.docx"
    word.save(path)
    content = path.read_bytes()
    text, _ = _parse_docx(str(path))
    job_id = uuid4()
    document = DocumentModel(
        document_id=job_id, source_name=path.name, file_type="docx",
        source_version=f"sha256:{hashlib.sha256(content).hexdigest()}",
        text=text, blocks=[], parser_name="test", parser_version="1",
    )
    storage = JobStorage(tmp_path / "jobs", 25 * 1024 * 1024)
    storage.save_bytes(job_id, path.name, content)

    for revision in (text, text.replace("Pro", "Basic")):
        layout = build_review_preview(document, storage, url, revision)
        covered = {i for page in layout.pages for g in page.glyphs for i in range(g.start, g.end)}
        assert {i for i, char in enumerate(revision) if not char.isspace()} <= covered
        assert layout.notice is None
    assert path.read_bytes() == content
