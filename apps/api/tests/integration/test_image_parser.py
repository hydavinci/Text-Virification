from __future__ import annotations

from hashlib import sha256
from pathlib import Path

import pymupdf
import pytest
from docx import Document

from text_verification.document_processing.errors import (
    OcrLayoutError,
    OcrOutputError,
    OcrProcessingError,
    OcrUnavailableError,
)
from text_verification.document_processing.ocr_provider import OcrTextBox, SupportedOcrLanguage
from text_verification.domain.documents import FileType
from text_verification.domain.ports import VerificationProgressStage
from text_verification.exporters.docx_reconstruction import DocxReconstructionExporter
from text_verification.parsers.errors import ParserError
from text_verification.parsers.image_parser import ImageParser, ImageResourceLimits


def _image(path: Path, file_type: FileType, *, width: int = 200, height: int = 120) -> Path:
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height), False)
    pixmap.clear_with(0xFFFFFF)
    path.write_bytes(pixmap.tobytes("png" if file_type is FileType.PNG else "jpeg"))
    return path


def _box(
    text: str,
    bbox: tuple[float, float, float, float],
    *,
    confidence: float = 0.99,
) -> OcrTextBox:
    x0, y0, x1, y1 = bbox
    return OcrTextBox(
        text=text,
        confidence=confidence,
        bbox=((x0, y0), (x1, y0), (x1, y1), (x0, y1)),
    )


class FakeOcr:
    def __init__(self, output: list[OcrTextBox]) -> None:
        self.output = output
        self.calls: list[tuple[object, str]] = []

    def recognize(self, image: object, language: str) -> list[OcrTextBox]:
        self.calls.append((image, language))
        return self.output


@pytest.mark.parametrize("file_type", [FileType.PNG, FileType.JPG])
def test_image_parser_emits_canonical_ocr_document(
    tmp_path: Path,
    file_type: FileType,
) -> None:
    source = _image(tmp_path / f"scan.{file_type.value}", file_type)
    ocr = FakeOcr(
        [
            _box("TITLE", (10, 10, 180, 35)),
            _box("body text", (10, 55, 180, 68)),
        ]
    )
    stages: list[VerificationProgressStage] = []

    document = ImageParser(file_type=file_type, ocr=ocr, ocr_language="en").parse_with_progress(
        source,
        progress_observer=stages.append,
    )

    assert document.source_version == f"sha256:{sha256(source.read_bytes()).hexdigest()}"
    assert document.file_type is file_type
    assert document.source_name == source.name
    assert document.text == "TITLE\nbody text"
    assert [block.kind for block in document.blocks] == ["heading", "paragraph"]
    assert all(block.source_locator["source"] == "ocr" for block in document.blocks)
    assert document.metadata.pdf is None
    assert stages == [VerificationProgressStage.OCR]
    assert [language for _, language in ocr.calls] == ["en"]


def test_image_parser_preserves_ocr_table_structure(tmp_path: Path) -> None:
    source = _image(tmp_path / "table.png", FileType.PNG)
    ocr = FakeOcr(
        [
            _box("A1", (10, 10, 40, 20)),
            _box("B1", (80, 10, 110, 20)),
            _box("A2", (10, 40, 40, 50)),
            _box("B2", (80, 40, 110, 50)),
        ]
    )

    document = ImageParser(file_type=FileType.PNG, ocr=ocr).parse(source)

    assert [block.kind for block in document.blocks] == ["table_cell"] * 4
    assert [
        (block.table_index, block.row_index, block.cell_index, block.text)
        for block in document.blocks
    ] == [
        (0, 0, 0, "A1"),
        (0, 0, 1, "B1"),
        (0, 1, 0, "A2"),
        (0, 1, 1, "B2"),
    ]


def test_image_parser_uses_grid_lines_to_preserve_empty_cells(tmp_path: Path) -> None:
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=200, height=160)
        for x in (20, 90, 160):
            page.draw_line((x, 20), (x, 120), width=1)
        for y in (20, 70, 120):
            page.draw_line((20, y), (160, y), width=1)
        source = tmp_path / "grid.png"
        page.get_pixmap().save(source)
    ocr = FakeOcr([
        _box("名称", (30, 35, 60, 48)),
        _box("说明", (30, 85, 60, 98)),
        _box("正文", (100, 85, 130, 98)),
    ])

    document = ImageParser(file_type=FileType.PNG, ocr=ocr).parse(source)

    assert [block.kind for block in document.blocks] == ["table_cell"] * 3
    assert [(block.row_index, block.cell_index) for block in document.blocks] == [
        (0, 0), (1, 0), (1, 1),
    ]
    assert all(block.source_locator["table_shape"] == {"rows": 2, "columns": 2}
               for block in document.blocks)
    assert all(document.text[block.global_start:block.global_end] == block.text
               for block in document.blocks)
    rebuilt = Document(DocxReconstructionExporter().export(document, tmp_path / "grid.docx"))
    assert len(rebuilt.tables) == 1
    assert [[cell.text for cell in row.cells] for row in rebuilt.tables[0].rows] == [
        ["名称", ""], ["说明", "正文"],
    ]


def test_image_parser_converts_pixel_font_height_using_embedded_dpi(tmp_path: Path) -> None:
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 300, 180), False)
    pixmap.clear_with(255)
    pixmap.set_dpi(300, 300)
    source = tmp_path / "300dpi.png"
    pixmap.save(source)
    document = ImageParser(
        file_type=FileType.PNG,
        ocr=FakeOcr([_box("正文", (20, 20, 200, 70))]),
    ).parse(source)

    assert document.blocks[0].style["font"] == {"size": pytest.approx(12.0, abs=0.05)}
    rebuilt = Document(DocxReconstructionExporter().export(document, tmp_path / "dpi.docx"))
    assert rebuilt.paragraphs[0].runs[0].font.size.pt == pytest.approx(12.0, abs=0.5)


def test_image_parser_rejects_blank_ocr_result(tmp_path: Path) -> None:
    source = _image(tmp_path / "blank.png", FileType.PNG)

    with pytest.raises(OcrProcessingError) as raised:
        ImageParser(file_type=FileType.PNG, ocr=FakeOcr([])).parse(source)

    assert raised.value.code == "ocr_no_text"
    assert raised.value.stage == "ocr"
    assert raised.value.retryable is False


def test_image_parser_reports_missing_ocr_engine(tmp_path: Path) -> None:
    source = _image(tmp_path / "scan.png", FileType.PNG)

    with pytest.raises(OcrUnavailableError):
        ImageParser(file_type=FileType.PNG, ocr=None).parse(source)


def test_image_parser_rejects_spoofed_corrupt_and_oversized_images(tmp_path: Path) -> None:
    spoofed = tmp_path / "spoofed.png"
    spoofed.write_bytes(b"\xff\xd8\xff\xe0not really png\xff\xd9")
    corrupt = tmp_path / "corrupt.png"
    corrupt.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 40)
    oversized = _image(tmp_path / "oversized.png", FileType.PNG, width=101, height=10)
    parser = ImageParser(
        file_type=FileType.PNG,
        ocr=FakeOcr([]),
        limits=ImageResourceLimits(max_width=100),
    )

    for source in (spoofed, corrupt, oversized):
        with pytest.raises(ParserError):
            parser.parse(source)


@pytest.mark.parametrize(
    ("language", "words", "second_y", "expected"),
    [
        ("zh", ("测", "试"), 35, "测试"),
        ("zh", ("测", "试"), 34, "测试"),
        ("ja", ("日", "本"), 34, "日本"),
        ("en", ("hello", "world"), 34, "hello world"),
    ],
)
def test_grid_cell_clusters_rows_before_language_aware_join(
    tmp_path: Path, language: SupportedOcrLanguage, words: tuple[str, str],
    second_y: int, expected: str,
) -> None:
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=200, height=160)
        for x in (20, 90, 160):
            page.draw_line((x, 20), (x, 120), width=1)
        for y in (20, 70, 120):
            page.draw_line((20, y), (160, y), width=1)
        source = tmp_path / "grid.png"
        page.get_pixmap().save(source)
    document = ImageParser(
        file_type=FileType.PNG,
        ocr=FakeOcr([
            _box(words[0], (30, 35, 48, 48)),
            _box(words[1], (50, second_y, 78, second_y + 13)),
            _box("X", (100, 85, 130, 98)),
        ]),
        ocr_language=language,
    ).parse(source)
    assert document.blocks[0].text == expected
    assert [box["text"] for box in document.blocks[0].source_locator["boxes"]] == list(words)
    assert all(document.text[b.global_start:b.global_end] == b.text for b in document.blocks)
    rebuilt = Document(DocxReconstructionExporter().export(document, tmp_path / "grid.docx"))
    assert [[cell.text for cell in row.cells] for row in rebuilt.tables[0].rows] == [
        [expected, ""], ["", "X"],
    ]


@pytest.mark.parametrize("reverse", [False, True])
@pytest.mark.parametrize("shift", [0.0, 0.1])
def test_image_ocr_discards_zero_confidence_and_coalesces_duplicates(
    tmp_path: Path, reverse: bool, shift: float,
) -> None:
    source = _image(tmp_path / "scan.png", FileType.PNG)
    boxes = [
        _box("Wrong", (10, 10, 80, 25), confidence=0.5),
        _box("Correct", (10 + shift, 10, 80 + shift, 25), confidence=0.99),
        _box("Invisible", (10, 50, 80, 65), confidence=0.0),
    ]
    document = ImageParser(
        file_type=FileType.PNG, ocr=FakeOcr(list(reversed(boxes)) if reverse else boxes),
    ).parse(source)
    assert document.text == "Correct"
    assert len(document.blocks[0].source_locator["boxes"]) == 1
    assert document.text[document.blocks[0].global_start:document.blocks[0].global_end] == "Correct"


@pytest.mark.parametrize(
    "box",
    [
        _box("Outside", (-1, 10, 80, 25)),
        _box("Outside", (10, 10, 201, 25)),
        OcrTextBox.model_construct(
            text="Invalid", confidence=float("nan"),
            bbox=((10, 10), (80, 10), (80, 25), (10, 25)),
        ),
    ],
)
def test_image_parser_rejects_invalid_provider_boxes(tmp_path: Path, box: OcrTextBox) -> None:
    source = _image(tmp_path / "scan.png", FileType.PNG)
    with pytest.raises(OcrOutputError):
        ImageParser(file_type=FileType.PNG, ocr=FakeOcr([box])).parse(source)


def test_image_parser_bounds_provider_text_characters(tmp_path: Path) -> None:
    source = _image(tmp_path / "scan.png", FileType.PNG)
    with pytest.raises(OcrLayoutError, match="max_ocr_text_chars"):
        ImageParser(
            file_type=FileType.PNG,
            ocr=FakeOcr([_box("too much text", (10, 10, 80, 25))]),
            limits=ImageResourceLimits(max_ocr_text_characters=5),
        ).parse(source)


def test_image_parser_bounds_duplicate_candidate_inspections(tmp_path: Path) -> None:
    source = _image(tmp_path / "scan.png", FileType.PNG)
    with pytest.raises(OcrLayoutError, match="max_ocr_duplicate_candidate_inspections"):
        ImageParser(
            file_type=FileType.PNG,
            ocr=FakeOcr([
                _box("a", (10, 10, 80, 25)),
                _box("a", (10.1, 10, 80.1, 25)),
            ]),
            limits=ImageResourceLimits(max_ocr_candidate_checks=0),
        ).parse(source)
